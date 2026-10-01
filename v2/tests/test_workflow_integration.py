"""Production entry point with real workbook bytes and a simulated upload boundary."""

from hashlib import sha256
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from streamlit.delta_generator import DeltaGenerator
from streamlit.testing.v1 import AppTest

from test_source_parsers import adcom_source, student_source


def test_import_edit_solve_download_and_failed_reimport(monkeypatch):
    # AppTest cannot drive file upload controls. Keep their real rendering and
    # callbacks, and supply bytes at this one boundary; all downstream code runs.
    files = {
        "Student interviewer availability": student_source(),
        "Adcom interviewer availability": adcom_source(),
    }
    original_upload = DeltaGenerator.file_uploader

    def uploaded(self, label, *args, **kwargs):
        value = original_upload(self, label, *args, **kwargs)
        return BytesIO(files[label]) if label in files else value

    monkeypatch.setattr(DeltaGenerator, "file_uploader", uploaded)
    root_app = Path(__file__).resolve().parents[2] / "app.py"
    app = AppTest.from_file(str(root_app), default_timeout=20).run()
    next(item for item in app.button if item.label == "Find interview periods and continue").click().run()
    assert not app.exception
    assert len(app.session_state["v2_import"].problem.slots) == 2
    app.number_input(key="v2_bulk_capacity").set_value(1).run()
    app.button(key="v2_apply_bulk_counts").click().run()
    app.checkbox(key="v2_recommended_rules").uncheck().run()
    app.number_input(key="student_min").set_value(0)
    app.number_input(key="adcom_min").set_value(0)
    app.run()
    next(item for item in app.button if item.label == "Create schedule using standard rules").click().run()
    assert not app.exception
    result = app.session_state["v2_result"]
    assert result.succeeded
    assert result.settings["source_student_sha256"] == sha256(files["Student interviewer availability"]).hexdigest()
    assert result.settings["source_adcom_sha256"] == sha256(files["Adcom interviewer availability"]).hexdigest()
    full = load_workbook(BytesIO(app.session_state["v2_report"]), data_only=True)
    simple = load_workbook(BytesIO(app.session_state["v2_simplified_report"]), data_only=True)
    assert len(full.sheetnames) == 9
    rows = list(full["Assignments"].values)
    assert len(rows) - 1 == len(result.assignments)
    assigned_names = [dict(zip(rows[0], row))["Interviewer Name"] for row in rows[1:]]
    simplified_names = [value for row in list(simple["Schedule"].values)[1:]
                        for value in row[2:] if value is not None]
    assert sorted(assigned_names) == sorted(simplified_names)
    downloads = [item for item in app.get("download_button") if "schedule" in item.label.lower()]
    assert len(downloads) == 2
    assert all(not item.proto.disabled for item in downloads)

    # Even a failed replacement import must not leave the prior run usable.
    files["Student interviewer availability"] = b"not an Excel file"
    next(item for item in app.button if item.label == "Find interview periods and continue").click().run()
    assert not app.exception
    assert "v2_result" not in app.session_state
    assert "v2_report" not in app.session_state
    assert "v2_import" not in app.session_state
    assert any("files could not be checked" in item.value for item in app.error)
