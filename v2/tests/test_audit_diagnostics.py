from dataclasses import replace
from io import BytesIO
import json

from openpyxl import load_workbook
import pytest

from interview_scheduler_v2 import (
    GroupPolicy, Interviewer, InterviewerGroup, LockedAssignment,
    RelaxationMode, SchedulingProblem, validate_problem,
)
from interview_scheduler_v2.audit import fingerprint
from interview_scheduler_v2.optimization import solve
from interview_scheduler_v2.presentation import present_diagnostics
from interview_scheduler_v2.reporting import build_workbook
from test_solver import config, make_slot


def test_minimum_failure_keeps_name_group_path_and_limiting_counts():
    person = Interviewer.create(name="Alex Unavailable", group=InterviewerGroup.STUDENT)
    problem = SchedulingProblem((person,), (make_slot("s1", 8, capacity=3),))
    cfg = config(student=GroupPolicy(3, 3, 5, 2))
    issue = validate_problem(problem, cfg).by_code("MIN_TOTAL_INFEASIBLE")[0]
    result = solve(problem, scenario="Failure", config=cfg)
    diagnostic = next(item for item in result.diagnostics if item.code == issue.code)
    assert diagnostic.interviewer_name == person.name
    assert diagnostic.interviewer_id == person.id
    assert diagnostic.group == person.group
    assert diagnostic.path == "interviewers[0]"
    assert diagnostic.context == issue.context
    assert (diagnostic.expected, diagnostic.actual) == (3, 0)
    message = next(item for item in present_diagnostics(result.diagnostics) if item.key == issue.code)
    assert any("Alex Unavailable" in detail and "at most 0" in detail for detail in message.details)
    workbook = load_workbook(BytesIO(build_workbook(result, problem, cfg)))
    rows = list(workbook["Constraint_Diagnostics"].values)
    row = next(dict(zip(rows[0], values)) for values in rows[1:] if values[1] == issue.code)
    assert row["Source"] == issue.path
    assert row["Interviewer Name"] == person.name
    assert (row["Expected"], row["Actual"]) == (3, 0)


def test_locked_daily_conflict_identifies_person_date_and_counts():
    slots = (make_slot("s1", 8), make_slot("s2", 10))
    person = Interviewer.create(name="Blair Locked", group=InterviewerGroup.STUDENT,
                                available_slot_ids=[slot.id for slot in slots])
    locks = tuple(LockedAssignment(person.id, slot.id, slot.local_date) for slot in slots)
    problem = SchedulingProblem((person,), slots, locks)
    result = solve(problem, scenario="Locks", config=config(student=GroupPolicy(0, 1, 3, 1)))
    diagnostic = next(item for item in result.diagnostics if item.code == "LOCKS_EXCEED_MAX_PER_DAY")
    assert diagnostic.interviewer_name == person.name
    assert diagnostic.assignment_date == slots[0].local_date
    assert (diagnostic.expected, diagnostic.actual) == (1, 2)


def test_fingerprints_are_stable_and_sensitive_to_effective_inputs_and_policies():
    slots = (make_slot("s1", 8), make_slot("s2", 10))
    person = Interviewer.create(name="One", group=InterviewerGroup.STUDENT,
                                available_slot_ids=["s2", "s1"])
    problem = SchedulingProblem((person,), slots)
    same = replace(problem, interviewers=(replace(person, available_slot_ids=frozenset(["s1", "s2"])),))
    assert fingerprint(problem) == fingerprint(same)
    assert fingerprint(problem) != fingerprint(replace(problem, slots=slots[::-1]))
    assert fingerprint(problem) != fingerprint(replace(problem, interviewers=(replace(person, historical_prior_count=1),)))
    assert fingerprint(problem) != fingerprint(replace(problem, interviewers=(replace(person, preference_by_slot={"s1": 5}),)))
    assert fingerprint(problem) != fingerprint(replace(problem, locked_assignments=(LockedAssignment(person.id, "s1", slots[0].local_date),)))
    assert fingerprint(config()) != fingerprint(replace(config(), person_policies={person.id: GroupPolicy(0, 1, 2, 1, 1)}))


def test_workbook_records_versions_hashes_run_time_stages_and_full_overrides():
    slot = make_slot("s1", 8)
    person = Interviewer.create(name="One", group=InterviewerGroup.STUDENT, available_slot_ids=[slot.id])
    problem = SchedulingProblem((person,), (slot,))
    policy = GroupPolicy(0, 1, 2, 1, 1)
    cfg = replace(config(), person_policies={person.id: policy})
    result = solve(problem, scenario="Audit", config=cfg)
    again = solve(problem, scenario="Audit", config=cfg)
    assert result.settings["input_sha256"] == again.settings["input_sha256"]
    assert result.settings["configuration_sha256"] == again.settings["configuration_sha256"]
    assert result.settings["run_id"] != again.settings["run_id"]
    relaxed = solve(problem, scenario="Audit", config=cfg, relaxation_mode=RelaxationMode.MINIMUMS)
    assert result.settings["configuration_sha256"] != relaxed.settings["configuration_sha256"]
    workbook = load_workbook(BytesIO(build_workbook(result, problem, cfg)), data_only=True)
    settings = dict(list(workbook["Run_Settings"].values)[1:])
    assert settings["Generated At ET"] == result.settings["generated_at_et"]
    for key in ("application_version", "package_source_sha256", "schema_version",
                "input_sha256", "configuration_sha256", "python_version", "version:ortools"):
        assert settings[f"Setting: {key}"] == result.settings[key]
    override = json.loads(settings[f"Setting: person_policy:{person.id}"])
    assert override["max_per_day"] == override["min_per_active_day"] == 1
    assert json.loads(settings["Setting: solver_stages_json"])[-1]["status"] == "OPTIMAL"
    rows = list(workbook["Interviewer_Summary"].values)
    summary = dict(zip(rows[0], rows[1]))
    assert summary["Maximum Per Day"] == summary["Minimum Per Active Day"] == 1


@pytest.mark.parametrize("changed", ["input", "config"])
def test_audit_export_rejects_inputs_or_settings_from_another_run(changed):
    slot = make_slot("s1", 8)
    person = Interviewer.create(name="One", group=InterviewerGroup.STUDENT,
                                available_slot_ids=[slot.id])
    problem = SchedulingProblem((person,), (slot,))
    cfg = config()
    result = solve(problem, scenario="Original", config=cfg)
    if changed == "input":
        problem = replace(problem, slots=(replace(slot, capacity=2),))
    else:
        cfg = replace(cfg, student_priority_weight=6)
    with pytest.raises(ValueError, match="differ"):
        build_workbook(result, problem, cfg)
